"""Institutional exports reuse the report query; no additional visibility scope."""
import io
from app.core.branding import PRIMARY_HEX
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from app.modules.correspondence.insights import BOLIVIA
from app.modules.correspondence.pdf.report_pdf import generate_report_pdf

INSTITUTION = "Gobierno Autónomo Municipal de Sipe Sipe"
KINDS = {"registered": "Registrada", "received": "Recibida", "sent": "Enviada", "concluded": "Concluida"}
HEADERS = ["Hoja de Ruta", "Fecha (Bolivia)", "Tipo", "Asunto / Referencia", "Origen", "Destino actual", "Estado"]


def export_filename(extension, selection):
    kind = {"registered": "correspondencia", "received": "recibidos", "sent": "enviados", "concluded": "concluidos"}[selection.get("kind", "registered")]
    dates = "_".join(str(selection[key]) for key in ("date_from", "date_to") if selection.get(key))
    suffix = dates or datetime.now(BOLIVIA).strftime("%Y-%m-%d")
    return re.sub(r"[^a-zA-Z0-9_.-]", "_", f"reporte_{kind}_{suffix}.{extension}")


def report_content(service, user, selection):
    # Validate permissions' scope/filter values using exactly the existing report query.
    query, event_date = service.query(user, **selection)
    summary = service.report(user, page=1, page_size=1, **selection)["summary"]
    options = service.options(user)
    applied = []
    for key, label, source in [("document_type_id", "Tipo documental", "types"), ("unit_id", "Unidad actual", "units"), ("employee_id", "Funcionario responsable", "employees")]:
        if selection.get(key):
            name = next((item["name"] for item in options[source] if str(item["id"]) == str(selection[key])), "Sin resultados")
            applied.append(f"{label}: {name}")
    if selection.get("status"):
        applied.append("Estado: " + ("Pendiente" if selection["status"] == "ACTIVE" else "Concluido"))
    if selection.get("search") and selection["search"].strip():
        applied.append("Búsqueda: " + selection["search"].strip())
    metadata = [
        ("Tipo de reporte", KINDS[selection.get("kind", "registered")]),
        ("Periodo", f"{selection.get('date_from') or 'Sin inicio'} / {selection.get('date_to') or 'Sin fin'}"),
        ("Filtros aplicados", " · ".join(applied) or "Sin filtros adicionales"),
        ("Generado (Bolivia)", datetime.now(BOLIVIA).strftime("%d/%m/%Y %H:%M")),
    ]
    def rows():
        for row in service.db.execute(service.rows_query(query, event_date).execution_options(yield_per=200)):
            yield [row.route_number, row.date.astimezone(BOLIVIA).strftime("%d/%m/%Y %H:%M") if row.date else "",
                row.document_type, " / ".join(str(v) for v in (row.subject, row.reference) if v), row.origin or "—",
                " · ".join(str(v) for v in (row.destination, row.employee) if v) or "Sin destino",
                "Pendiente" if row.status == "ACTIVE" else "Concluido"]
    return metadata, summary, rows()


def generate_xlsx(metadata, summary, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Correspondencia"
    sheet.sheet_view.showGridLines = False
    def append(values):
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                # User-controlled text is always a string, never a formula or hyperlink.
                cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11, color="172A3A")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    append([INSTITUTION]); sheet.merge_cells("A1:G1")
    sheet['A1'].font = Font(name="Calibri", size=15, bold=True, color=PRIMARY_HEX)
    append(["Reporte de Correspondencia"]); sheet.merge_cells("A2:G2")
    for label, value in metadata:
        append([label, value]); sheet.merge_cells(start_row=sheet.max_row, start_column=2, end_row=sheet.max_row, end_column=7)
        sheet.row_dimensions[sheet.max_row].height = max(20, 16 * ((len(value) // 115) + 1))
    append([])
    append(["Total", summary["total"], "Pendientes", summary["pending"], "Concluidos", summary["concluded"]])
    append([])
    append(HEADERS)
    header = sheet.max_row
    for cell in sheet[header]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY_HEX)
    widths = [24, 22, 24, 65, 35, 48, 18]
    for index, width in enumerate(widths, 1):sheet.column_dimensions[get_column_letter(index)].width = width
    for values in rows:
        append(values)
        sheet.row_dimensions[sheet.max_row].height = min(409, max(30, 16 * max((len(str(v)) // widths[i] + str(v).count("\n") + 1) for i, v in enumerate(values))))
    sheet.freeze_panes = f"A{header+1}"
    sheet.auto_filter.ref = f"A{header}:G{sheet.max_row}"
    sheet.print_title_rows = f"{header}:{header}"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    stream = io.BytesIO();workbook.save(stream)
    return stream.getvalue()


def generate_export(service, user, extension, selection):
    metadata, summary, rows = report_content(service, user, selection)
    if extension == "xlsx":return generate_xlsx(metadata, summary, rows)
    return generate_report_pdf(metadata, summary, rows)
