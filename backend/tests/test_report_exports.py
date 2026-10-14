import csv
import io
import uuid
from datetime import datetime
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from app.modules.correspondence.report_exports import generate_xlsx
from app.modules.correspondence.pdf.report_pdf import generate_report_pdf
from tests.test_final_blocks import document, headers, reader, grant

FORMATS = ["csv", "xlsx", "pdf"]


def text(response, fmt):
    assert response.status_code == 200, response.text[:300] if response.status_code != 200 else ""
    if fmt == "csv":return response.content.decode("utf-8-sig")
    if fmt == "xlsx":
        assert response.content[:2] == b"PK"
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Correspondencia"]
        assert book.active.freeze_panes == "A11"
        return " ".join(str(c.value or "") for row in book.active for c in row)
    assert response.content.startswith(b"%PDF")
    return " ".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)


@pytest.mark.parametrize("fmt", FORMATS)
def test_export_real_format_filters_and_personal_scope(client, db, document, second_user, fmt):
    grant(db, second_user)
    params = {"kind": "received", "status": "ACTIVE", "search": document["route_number"],
        "date_from": "2000-01-01", "date_to": "2099-12-31"}
    exported = client.get(f"/reports/export.{fmt}", params=params, headers=headers(second_user))
    content = text(exported, fmt)
    assert document["route_number"] in content and "Final QA visible" in content
    assert "attachment;" in exported.headers["content-disposition"]
    assert f"reporte_recibidos_2000-01-01_2099-12-31.{fmt}" in exported.headers["content-disposition"]
    assert "application/" in exported.headers["content-type"] if fmt != 'csv' else "text/csv" in exported.headers["content-type"]
    if fmt != "csv":
        assert "Gobierno Autónomo Municipal de Sipe Sipe" in content or "GOBIERNO AUTÓNOMO MUNICIPAL DE SIPE SIPE" in content
        assert "Filtros aplicados" in content
        assert document["id"] not in content
    params['search'] = 'DOES_NOT_EXIST_QA'
    assert document['route_number'] not in text(client.get(f'/reports/export.{fmt}', params=params, headers=headers(second_user)),fmt)
    params.pop('search');params['date_to']='2000-01-02'
    assert document['route_number'] not in text(client.get(f'/reports/export.{fmt}', params=params, headers=headers(second_user)),fmt)


@pytest.mark.parametrize("fmt", FORMATS)
def test_export_unit_scope_and_employee_filter(client, db, document, second_org_unit, second_user, fmt):
    scoped, _, _ = reader(db, second_org_unit);grant(db, scoped)
    params = {'unit_id': str(second_org_unit.id), 'employee_id': str(second_user.employee_id), 'status':'ACTIVE'}
    assert document['route_number'] in text(client.get(f'/reports/export.{fmt}',params=params,headers=headers(scoped)),fmt)
    params['status']='CONCLUDED'
    assert document['route_number'] not in text(client.get(f'/reports/export.{fmt}',params=params,headers=headers(scoped)),fmt)


@pytest.mark.parametrize("fmt", FORMATS)
def test_export_other_unit_rejected(client, db, document, second_user, org_unit, fmt):
    grant(db,second_user)
    assert client.get(f'/reports/export.{fmt}',params={'unit_id':str(org_unit.id)},headers=headers(second_user)).status_code == 403


@pytest.mark.parametrize("fmt", FORMATS)
def test_export_requires_reports_permission(client, document, second_user, fmt):
    assert client.get(f'/reports/export.{fmt}',headers=headers(second_user)).status_code == 403
    assert client.get(f'/reports/export.{fmt}').status_code == 401


@pytest.mark.parametrize("fmt", FORMATS)
def test_export_outsider_has_no_private_rows(client, db, document, org_unit, fmt):
    outsider, _, _=reader(db,org_unit);grant(db,outsider)
    assert document['route_number'] not in text(client.get(f'/reports/export.{fmt}',headers=headers(outsider)),fmt)


@pytest.mark.parametrize("fmt", FORMATS)
def test_export_same_document_type_and_reversed_dates(client, db, document, second_user, fmt):
    grant(db,second_user)
    assert document['route_number'] not in text(client.get(f'/reports/export.{fmt}',params={'document_type_id':str(uuid.uuid4())},headers=headers(second_user)),fmt)
    assert client.get(f'/reports/export.{fmt}',params={'date_from':'2026-10-31','date_to':'2026-10-01'},headers=headers(second_user)).status_code==422


def test_xlsx_formula_text_and_metadata_are_not_executable():
    book=load_workbook(io.BytesIO(generate_xlsx([('Búsqueda','=HYPERLINK("x")')],{'total':1,'pending':1,'concluded':0},[['HR-1','14/10/2026','NOTA','=1+1','+1','@SUM(A1:A2)','Pendiente']])))
    for row in book.active:
        for cell in row:
            assert cell.data_type != 'f' and cell.hyperlink is None
    assert any(cell.value=='=1+1' for row in book.active for cell in row)


def test_pdf_pagination_long_text_and_repeated_headers():
    rows=[['HR-'+str(i),'14/10/2026','NOTA','Texto administrativo '+('muy largo & <seguro> '* (600 if i==0 else 5)), 'Origen QA','Unidad QA','Pendiente'] for i in range(55)]
    raw=generate_report_pdf([('Periodo','Octubre 2026')],{'total':55,'pending':55,'concluded':0},rows)
    pages=PdfReader(io.BytesIO(raw)).pages
    assert len(pages)>2
    content=' '.join(p.extract_text() for p in pages)
    assert 'HR-54' in content and '<seguro>' in content
    for i,page in enumerate(pages,1):
        page_text=page.extract_text()
        assert 'REPORTE DE CORRESPONDENCIA' in page_text
        assert 'Asunto / Referencia' in page_text
        assert f'Página {i} de {len(pages)}' in page_text
